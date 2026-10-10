"""PDF and Excel export stamped with filters, user and time (FR-65, UC-28)."""
import io

from django.utils import timezone


def _stamp(title, filters_text, user):
    return [title, f"Filters: {filters_text}",
            f"Exported by {user.get_full_name() or user.email} on "
            f"{timezone.localtime():%d %b %Y %H:%M} (Cairo)"]


def to_excel(title, filters_text, user, kpis, tables):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    for i, line in enumerate(_stamp(title, filters_text, user), start=1):
        ws.cell(row=i, column=1, value=line).font = Font(bold=(i == 1), size=14 if i == 1 else 10)
    row = 5
    for label, value in kpis:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=str(value))
        row += 1
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 22
    header_fill = PatternFill("solid", fgColor="DDE7F0")
    for n, t in enumerate(tables, start=1):
        sheet = wb.create_sheet(f"{n}. {t['title']}"[:31].replace("/", "-"))
        for i, line in enumerate(_stamp(t["title"], filters_text, user), start=1):
            sheet.cell(row=i, column=1, value=line).font = Font(bold=(i == 1))
        for c, h in enumerate(t["headers"], start=1):
            cell = sheet.cell(row=5, column=c, value=h)
            cell.font = Font(bold=True)
            cell.fill = header_fill
            sheet.column_dimensions[cell.column_letter].width = max(14, len(str(h)) + 4)
        for r, data in enumerate(t["rows"], start=6):
            for c, v in enumerate(data, start=1):
                sheet.cell(row=r, column=c, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(title, filters_text, user, kpis, tables):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm,
                            rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
                            title=title)
    styles = getSampleStyleSheet()
    story = []
    stamp = _stamp(title, filters_text, user)
    story.append(Paragraph(stamp[0], styles["Title"]))
    story.append(Paragraph(stamp[1], styles["Normal"]))
    story.append(Paragraph(stamp[2], styles["Normal"]))
    story.append(Spacer(1, 6 * mm))
    if kpis:
        kt = Table([[k, str(v)] for k, v in kpis], hAlign="LEFT")
        kt.setStyle(TableStyle([("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                                ("GRID", (0, 0), (-1, -1), 0.3, colors.grey)]))
        story += [kt, Spacer(1, 6 * mm)]
    cell = styles["BodyText"]
    cell.fontSize = 8
    cell.leading = 10
    for t in tables:
        story.append(Paragraph(t["title"], styles["Heading3"]))
        if not t["rows"]:
            story.append(Paragraph("No data for the selected filters.", styles["Italic"]))
            continue
        data = [[Paragraph(f"<b>{h}</b>", cell) for h in t["headers"]]]
        data += [[Paragraph(str(v), cell) for v in r] for r in t["rows"]]
        tb = Table(data, repeatRows=1, hAlign="LEFT")
        tb.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDE7F0")),
                                ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                                ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story += [tb, Spacer(1, 5 * mm)]
    doc.build(story)
    return buf.getvalue()
