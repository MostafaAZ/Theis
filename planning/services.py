"""Data import, demand forecast and purchase plan (UC-23 to UC-26)."""
import csv
import io
import math
import re
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Sum
from django.utils import timezone

from core.models import AuditLog, BinLocation, ProductCategory, Region, Role, Supplier
from core.rules import rule, rule_int
from core.services import audit, notify, notify_roles
from inventory.models import (Batch, BatchStatus, BatchStock, Customer, Product,
                              PurchaseOrder, PurchaseOrderLine, StockMovement)
from inventory.services import evaluate_batch, free_qty, next_number

from . import forecasting
from .models import (DataImport, Forecast, ForecastLine, ForecastProduct, PurchasePlan,
                     PurchasePlanLine, SalesRecord)

# --------------------------------------------------------------------------
# UC-23: Import sales and stock data
# --------------------------------------------------------------------------
HEADER_ALIASES = {
    "order_number": ["order number", "order no", "order", "order_number", "so number"],
    "region": ["area", "region", "governorate", "zone"],
    "customer": ["customer", "customer name", "client", "partner"],
    "category": ["category", "product category"],
    "product": ["product", "product name", "item", "sku", "product code", "internal reference"],
    "quantity": ["quantity", "qty", "units", "quantity delivered", "on hand", "qty on hand"],
    "price": ["price", "unit price", "sale price"],
    "total": ["total", "subtotal", "amount"],
    "month": ["month"],
    "year": ["year"],
    "date": ["date", "order date", "invoice date"],
    "lot": ["lot", "lot number", "lot/serial number", "batch"],
    "expiry": ["expiry", "expiry date", "expiration date", "exp date"],
    "manufacture": ["manufacture date", "mfg date", "production date"],
    "location": ["location", "bin", "bin location", "location code"],
    "po_number": ["po number", "purchase order", "po", "reference"],
    "supplier": ["supplier", "vendor"],
    "expected": ["expected date", "scheduled date", "expected"],
    "name": ["name", "product name"],
    "unit": ["unit", "uom", "unit of measure"],
    "storage": ["storage", "storage condition"],
    "cost": ["cost", "unit cost"],
    "moq": ["moq", "minimum order quantity"],
    "lead_time": ["lead time", "lead time days", "lead_time_days"],
    "shelf_life": ["shelf life", "shelf life months"],
}

REQUIRED = {
    DataImport.SALES: ["product", "quantity", "region"],
    DataImport.STOCK: ["product", "lot", "expiry", "quantity", "location"],
    DataImport.OPEN_PO: ["po_number", "supplier", "product", "quantity"],
    DataImport.PRODUCTS: ["product", "name", "category", "price", "cost"],
}

MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


def read_table(uploaded):
    """Read an uploaded CSV or XLSX file into a header list and row dicts."""
    name = uploaded.name.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(uploaded, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return [], []
        headers = [str(h or "").strip() for h in rows[0]]
        data = [dict(zip(headers, r)) for r in rows[1:] if any(c not in (None, "") for c in r)]
        return headers, data
    if name.endswith(".csv"):
        text = uploaded.read().decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        headers = [h.strip() for h in (reader.fieldnames or [])]
        data = [{(k or "").strip(): v for k, v in r.items()} for r in reader]
        return headers, data
    raise ValidationError("Upload an Excel (.xlsx) or CSV file exported from Odoo.")


def map_columns(headers):
    lookup = {h.lower().strip(): h for h in headers}
    mapping = {}
    for field, aliases in HEADER_ALIASES.items():
        for a in aliases:
            if a in lookup:
                mapping[field] = lookup[a]
                break
    return mapping


def _norm(text):
    return re.sub(r"[^a-z0-9؀-ۿ]", "", str(text or "").lower().replace("&", "and"))


def region_index():
    idx = {}
    for r in Region.objects.filter(is_active=True):
        idx[_norm(r.name)] = r
        for a in (r.aliases or "").split("|"):
            if a.strip():
                idx[_norm(a)] = r
    return idx


def find_product(value, by_sku, by_name):
    v = str(value or "").strip()
    if not v:
        return None
    if v.upper() in by_sku:
        return by_sku[v.upper()]
    m = re.match(r"^\[(.+?)\]", v)  # Odoo "[SKU] Name" format
    if m and m.group(1).upper() in by_sku:
        return by_sku[m.group(1).upper()]
    return by_name.get(v.lower())


def _num(v):
    if v in (None, ""):
        return None
    try:
        return Decimal(str(v).replace(",", "").strip())
    except InvalidOperation:
        return None


def _date(v):
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%Y/%m/%d", "%b %Y", "%Y-%m"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _period(row, cols):
    if "month" in cols and "year" in cols:
        m, y = row.get(cols["month"]), row.get(cols["year"])
        try:
            month = MONTHS.get(str(m).strip()[:3].lower()) or int(float(m))
            return date(int(float(y)), month, 1)
        except (TypeError, ValueError):
            return None
    if "date" in cols:
        d = _date(row.get(cols["date"]))
        return d.replace(day=1) if d else None
    return None


@transaction.atomic
def stage_import(user, data_type, uploaded):
    """Validate an export and stage the rows for confirmation (UC-23 steps 2-5).

    Nothing is written to the operational tables until the officer confirms
    (NFR-13)."""
    headers, rows = read_table(uploaded)
    cols = map_columns(headers)
    missing = [c for c in REQUIRED[data_type] if c not in cols]
    if data_type == DataImport.SALES and not (("month" in cols and "year" in cols)
                                              or "date" in cols):
        missing.append("month and year (or date)")
    if missing:
        raise ValidationError("Required column missing: " + ", ".join(missing))
    by_sku = {p.sku.upper(): p for p in Product.objects.all()}
    by_name = {p.name.lower(): p for p in Product.objects.all()}
    regions = region_index()
    staged, messages, unknown_regions = [], [], set()
    accepted = corrected = rejected = 0
    for i, row in enumerate(rows, start=2):
        rec, problem, fixed = {"row": i}, None, False
        if data_type == DataImport.PRODUCTS:
            sku = str(row.get(cols["product"]) or "").strip()
            price, cost = _num(row.get(cols["price"])), _num(row.get(cols["cost"]))
            if not sku or price is None or cost is None:
                problem = "missing SKU, price or cost"
            else:
                rec.update(sku=sku, name=str(row.get(cols["name"]) or sku).strip(),
                           category=str(row.get(cols["category"]) or "").strip(),
                           price=str(price), cost=str(cost),
                           unit=str(row.get(cols.get("unit"), "") or "unit").strip(),
                           storage=str(row.get(cols.get("storage"), "") or "ambient").lower(),
                           moq=int(_num(row.get(cols.get("moq"))) or 1),
                           lead_time=int(_num(row.get(cols.get("lead_time"))) or 90),
                           shelf_life=int(_num(row.get(cols.get("shelf_life"))) or 24),
                           supplier=str(row.get(cols.get("supplier"), "") or "").strip())
        else:
            product = find_product(row.get(cols["product"]), by_sku, by_name)
            qty = _num(row.get(cols["quantity"]))
            if product is None:
                problem = f"unknown product '{row.get(cols['product'])}'"
            elif qty is None:
                problem = "missing quantity"
            else:
                rec.update(product=product.pk, quantity=int(qty))
        if not problem and data_type == DataImport.SALES:
            period = _period(row, cols)
            raw_region = str(row.get(cols["region"]) or "").strip()
            region = regions.get(_norm(raw_region))
            if period is None:
                problem = "missing date"
            elif not raw_region:
                problem = "missing region"
            else:
                price = _num(row.get(cols.get("price"))) if "price" in cols else None
                total = _num(row.get(cols.get("total"))) if "total" in cols else None
                prod = Product.objects.get(pk=rec["product"])
                price = price if price is not None else (
                    (total / rec["quantity"]) if total and rec["quantity"] else
                    prod.original_price)
                total = total if total is not None else price * rec["quantity"]
                rec.update(period=period.isoformat(), raw_region=raw_region,
                           region=region.pk if region else None,
                           customer=str(row.get(cols.get("customer"), "") or "").strip(),
                           order_number=str(row.get(cols.get("order_number"), "") or ""),
                           price=str(price), total=str(total))
                if region is None:
                    unknown_regions.add(raw_region)
                elif region.name != raw_region:
                    fixed = True
        if not problem and data_type == DataImport.STOCK:
            expiry = _date(row.get(cols["expiry"]))
            loc_code = str(row.get(cols["location"]) or "").strip()
            loc = BinLocation.objects.filter(code__iexact=loc_code).first()
            if expiry is None:
                problem = "missing or invalid expiry date"
            elif loc is None:
                problem = f"unknown location '{loc_code}'"
            else:
                mfg = _date(row.get(cols.get("manufacture"))) if "manufacture" in cols else None
                rec.update(lot=str(row.get(cols["lot"])).strip(), expiry=expiry.isoformat(),
                           manufacture=(mfg or expiry - timedelta(days=730)).isoformat(),
                           location=loc.pk)
        if not problem and data_type == DataImport.OPEN_PO:
            exp = _date(row.get(cols.get("expected"))) if "expected" in cols else None
            rec.update(po_number=str(row.get(cols["po_number"])).strip(),
                       supplier=str(row.get(cols["supplier"])).strip(),
                       expected=exp.isoformat() if exp else None)
        if problem:
            rejected += 1
            messages.append(f"Row {i}: rejected – {problem}.")
            continue
        if fixed:
            corrected += 1
        accepted += 1
        staged.append(rec)
    imp = DataImport.objects.create(
        data_type=data_type, file_name=uploaded.name[:200], uploaded_by=user,
        rows_total=len(rows), rows_accepted=accepted, rows_corrected=corrected,
        rows_rejected=rejected, messages=messages[:200], staged_rows=staged)
    return imp, sorted(unknown_regions)


@transaction.atomic
def confirm_import(user, imp, region_map=None):
    """UC-23 steps 6-7. ``region_map`` maps unrecognised region names to a
    reference region; the spelling is saved as an alias for next time."""
    if imp.status != DataImport.PREVIEW:
        raise ValidationError("This import has already been processed.")
    region_map = region_map or {}
    for raw, region in region_map.items():
        if region:
            aliases = [a for a in (region.aliases or "").split("|") if a]
            if raw not in aliases:
                region.aliases = "|".join(aliases + [raw])
                region.save(update_fields=["aliases"])
    rows = imp.staged_rows
    if imp.data_type == DataImport.SALES:
        unresolved = {r["raw_region"] for r in rows if not r.get("region")
                      and not region_map.get(r["raw_region"])}
        if unresolved:
            raise ValidationError("Map every unrecognised region before importing: "
                                  + ", ".join(sorted(unresolved)))
        customers = {c.name.lower(): c for c in Customer.objects.all()}
        objs = []
        for r in rows:
            region_id = r.get("region") or region_map[r["raw_region"]].pk
            cust = customers.get(r["customer"].lower()) if r["customer"] else None
            if r["customer"] and cust is None:
                cust = Customer.objects.create(name=r["customer"], region_id=region_id)
                customers[r["customer"].lower()] = cust
            objs.append(SalesRecord(
                order_number=r["order_number"], period=date.fromisoformat(r["period"]),
                region_id=region_id, customer=cust, customer_name=r["customer"],
                product_id=r["product"], quantity=r["quantity"], unit_price=Decimal(r["price"]),
                total=Decimal(r["total"]), data_import=imp))
        SalesRecord.objects.bulk_create(objs)
    elif imp.data_type == DataImport.PRODUCTS:
        for r in rows:
            cat, _ = ProductCategory.objects.get_or_create(name=r["category"] or "Uncategorised")
            sup = None
            if r["supplier"]:
                sup, _ = Supplier.objects.get_or_create(name=r["supplier"])
            storage = r["storage"] if r["storage"] in ("ambient", "refrigerated",
                                                       "frozen") else "ambient"
            Product.objects.update_or_create(sku=r["sku"], defaults=dict(
                name=r["name"], category=cat, original_price=Decimal(r["price"]),
                unit_cost=Decimal(r["cost"]), unit=r["unit"], storage_condition=storage,
                moq=r["moq"], lead_time_days=r["lead_time"], shelf_life_months=r["shelf_life"],
                supplier=sup))
    elif imp.data_type == DataImport.OPEN_PO:
        for r in rows:
            sup, _ = Supplier.objects.get_or_create(name=r["supplier"])
            po, _ = PurchaseOrder.objects.get_or_create(number=r["po_number"], defaults=dict(
                supplier=sup, expected_date=r["expected"] and date.fromisoformat(r["expected"])))
            line, created = PurchaseOrderLine.objects.get_or_create(
                order=po, product_id=r["product"], defaults={"quantity": r["quantity"]})
            if not created:
                line.quantity = r["quantity"]
                line.save()
    elif imp.data_type == DataImport.STOCK:
        for r in rows:
            product = Product.objects.get(pk=r["product"])
            loc = BinLocation.objects.get(pk=r["location"])
            batch, created = Batch.objects.get_or_create(
                product=product, lot_number=r["lot"], defaults=dict(
                    manufacture_date=date.fromisoformat(r["manufacture"]),
                    expiry_date=date.fromisoformat(r["expiry"]),
                    storage_condition=product.storage_condition, market=loc.warehouse.market,
                    quantity_received=r["quantity"], quantity_approved=r["quantity"],
                    status=BatchStatus.IN_STOCK, received_at=timezone.now(), received_by=user))
            stock, _ = BatchStock.objects.get_or_create(batch=batch, location=loc)
            delta = r["quantity"] - stock.quantity
            stock.quantity = r["quantity"]
            stock.save()
            batch.recompute_on_hand()
            batch.save()
            if delta:
                StockMovement.objects.create(batch=batch, movement_type=StockMovement.ADJUSTMENT,
                                             quantity=delta, to_location=loc, user=user,
                                             reference=f"Import #{imp.pk}")
            evaluate_batch(batch)
    imp.status = DataImport.IMPORTED
    imp.staged_rows = []
    imp.save()
    audit(user, AuditLog.DATA_SYNC, f"{imp.get_data_type_display()} imported from "
                                    f"{imp.file_name}", f"IMPORT-{imp.pk}",
          new=f"accepted={imp.rows_accepted} corrected={imp.rows_corrected} "
              f"rejected={imp.rows_rejected}")
    return imp


# --------------------------------------------------------------------------
# UC-24: Generate demand forecast
# --------------------------------------------------------------------------
def add_months(d, n):
    idx = d.year * 12 + d.month - 1 + n
    return date(idx // 12, idx % 12 + 1, 1)


def months_between(a, b):
    return (b.year - a.year) * 12 + b.month - a.month


def monthly_series(product, region, end_period):
    """Monthly quantities from the first sale up to ``end_period`` (inclusive),
    with zero for months without sales."""
    rows = (SalesRecord.objects.filter(product=product, region=region, period__lte=end_period)
            .values("period").annotate(q=Sum("quantity")).order_by("period"))
    data = {r["period"]: float(r["q"]) for r in rows}
    if not data:
        return []
    start = min(data)
    n = months_between(start, end_period) + 1
    return [data.get(add_months(start, i), 0.0) for i in range(n)]


def latest_sales_period():
    return SalesRecord.objects.order_by("-period").values_list("period", flat=True).first()


@transaction.atomic
def generate_forecast(user, horizon, products, category=None):
    """FR-43, FR-44 and UC-24 A1/A2."""
    end = latest_sales_period()
    if end is None:
        raise ValidationError("Import sales history before generating a forecast.")
    fc = Forecast.objects.create(created_by=user, horizon_months=horizon,
                                 start_period=add_months(end, 1), category=category)
    threshold = float(rule("mape_low_confidence_pct"))
    pairs = (SalesRecord.objects.filter(product__in=products).values_list(
        "product_id", "region_id").distinct())
    prods = {p.pk: p for p in products}
    regions = {r.pk: r for r in Region.objects.all()}
    for pid, rid in sorted(set(pairs)):
        series = monthly_series(prods[pid], regions[rid], end)
        method, values = forecasting.forecast(series, horizon)
        err = forecasting.backtest_mape(series)
        fp = ForecastProduct.objects.create(
            forecast=fc, product=prods[pid], region=regions[rid], method=method,
            history_months=len(series), mape=None if err is None else Decimal(f"{err:.2f}"),
            low_confidence=err is not None and err > threshold,
            insufficient_history=method == ForecastProduct.MANUAL)
        ForecastLine.objects.bulk_create([
            ForecastLine(forecast_product=fp, period=add_months(fc.start_period, i),
                         quantity=Decimal(f"{v:.2f}")) for i, v in enumerate(values)])
    return fc


@transaction.atomic
def set_manual_estimate(fp, monthly_qty, reason):
    """UC-24 A1: manual estimate for products with too little history."""
    if not reason:
        raise ValidationError("A reason is required for a manual estimate.")
    fp.manual_reason = reason
    fp.method = ForecastProduct.MANUAL
    fp.save()
    fp.lines.update(quantity=Decimal(monthly_qty))


def actuals_for(product, region, end, months=12):
    start = add_months(end, -(months - 1))
    rows = (SalesRecord.objects.filter(product=product, region=region, period__gte=start,
                                       period__lte=end)
            .values("period").annotate(q=Sum("quantity")))
    data = {r["period"]: r["q"] for r in rows}
    return [(add_months(start, i), data.get(add_months(start, i), 0)) for i in range(months)]


# --------------------------------------------------------------------------
# UC-25: Build and submit purchase plan
# --------------------------------------------------------------------------
def usable_stock(product, monthly_rate):
    """FR-46: stock projected to sell before it reaches Near-Expiry.

    Batches are consumed in FEFO order at the forecast rate; whatever a
    batch cannot sell before its near-expiry date is excluded."""
    ne_days = rule_int("near_expiry_days")
    today = timezone.localdate()
    daily = monthly_rate / 30.4
    batches = Batch.objects.filter(product=product, qty_on_hand__gt=0,
                                   status=BatchStatus.IN_STOCK).order_by("expiry_date")
    usable = excluded = 0
    consumed = 0.0
    for b in batches:
        qty = free_qty(b)
        if qty <= 0:
            continue
        if daily <= 0:
            excluded += qty
            continue
        deadline = max(0, (b.expiry_date - today).days - ne_days)
        sellable = int(max(0, min(qty, daily * deadline - consumed)))
        usable += sellable
        excluded += qty - sellable
        consumed += sellable
    return usable, excluded


def plan_line_figures(product, monthly_forecast):
    """FR-45 to FR-47 for one product.

    ``monthly_forecast`` is the forecast demand per month from the start of
    the planning period."""
    horizon = len(monthly_forecast)
    avg = sum(monthly_forecast) / horizon if horizon else 0.0
    lt_months = product.lead_time_days / 30.4
    window = lt_months + rule_int("coverage_months")
    whole = int(window)
    demand = sum(monthly_forecast[:min(whole, horizon)]) + avg * max(0, whole - horizon)
    demand += avg * (window - whole)
    safety = avg * rule_int("safety_stock_days") / 30.4
    usable, excluded = usable_stock(product, avg)
    open_po = sum(line.open_quantity for line in PurchaseOrderLine.objects.filter(
        product=product, order__status=PurchaseOrder.OPEN))
    recommended = max(0, math.ceil(demand + safety - usable - open_po))
    rounded = max(recommended, product.moq) if recommended > 0 else 0
    moq_surplus = rounded - recommended
    # Stock on arrival = what remains of current stock after the lead time,
    # plus the order; whatever cannot sell within the shelf life is left over.
    remaining_on_arrival = max(0.0, usable + open_po - avg * lt_months)
    left = max(0, math.ceil(rounded + remaining_on_arrival - avg * product.shelf_life_months))
    left = min(left, rounded)
    return dict(forecast_qty=math.ceil(demand), safety_stock=math.ceil(safety),
                usable_stock=usable, excluded_stock=excluded, open_po_qty=open_po,
                recommended_qty=recommended, moq=product.moq, rounded_qty=rounded,
                moq_surplus_qty=moq_surplus,
                moq_surplus_value=(product.unit_cost * moq_surplus).quantize(Decimal("0.01")),
                projected_left_at_expiry=left, final_qty=rounded, at_risk=left > 0)


@transaction.atomic
def build_plan(user, forecast, period_label):
    """UC-25 steps 1-5."""
    if not forecast.products.exists():
        raise ValidationError("The forecast has no products.")
    by_product = defaultdict(lambda: [0.0] * forecast.horizon_months)
    for fp in forecast.products.prefetch_related("lines"):
        for i, line in enumerate(fp.lines.all()):
            by_product[fp.product_id][i] += float(line.quantity)
    plan = PurchasePlan.objects.create(number=next_number("PP-", PurchasePlan),
                                       forecast=forecast, period_label=period_label,
                                       prepared_by=user)
    for product in Product.objects.filter(pk__in=by_product.keys()):
        PurchasePlanLine.objects.create(plan=plan, product=product,
                                        **plan_line_figures(product, by_product[product.pk]))
    return plan


@transaction.atomic
def override_line(line, final_qty, reason):
    """FR-48: overrides need a reason."""
    if line.plan.status not in (PurchasePlan.DRAFT, PurchasePlan.RETURNED):
        raise ValidationError("The plan is locked.")
    if line.plan.status == PurchasePlan.RETURNED and line.status != PurchasePlanLine.RETURNED:
        raise ValidationError("Only returned lines can be changed.")
    if final_qty < 0:
        raise ValidationError("Quantity cannot be negative.")
    if final_qty != line.rounded_qty and not reason:
        raise ValidationError("A reason is required to override the recommended quantity.")
    line.final_qty = final_qty
    line.override_reason = reason if final_qty != line.rounded_qty else ""
    line.save()


@transaction.atomic
def submit_plan(user, plan):
    if plan.status not in (PurchasePlan.DRAFT, PurchasePlan.RETURNED):
        raise ValidationError("Only draft or returned plans can be submitted.")
    for line in plan.lines.all():
        if line.final_qty != line.rounded_qty and not line.override_reason:
            raise ValidationError(f"Override on {line.product.name} needs a reason.")
    plan.lines.filter(status=PurchasePlanLine.RETURNED).update(status=PurchasePlanLine.PENDING,
                                                                gm_comment="")
    plan.status = PurchasePlan.SUBMITTED
    plan.submitted_at = timezone.now()
    plan.save()
    notify_roles(Role.GM, f"Purchase plan {plan.number} submitted for approval.",
                 f"/plans/{plan.pk}/")


@transaction.atomic
def decide_plan(user, plan, decisions, overall_comment=""):
    """UC-26 / FR-49: approve, return or reject line by line.

    ``decisions`` maps line id to ``(status, comment)``."""
    if plan.status != PurchasePlan.SUBMITTED:
        raise ValidationError("Only submitted plans can be decided.")
    lines = list(plan.lines.all())
    for line in lines:
        status, comment = decisions.get(line.pk, (PurchasePlanLine.APPROVED, ""))
        if status in (PurchasePlanLine.RETURNED, PurchasePlanLine.REJECTED) and not comment:
            raise ValidationError(f"Add a comment for {line.product.name}.")
        line.status = status
        line.gm_comment = comment
        line.save()
    statuses = {line.status for line in lines}
    if PurchasePlanLine.RETURNED in statuses:
        plan.status = PurchasePlan.RETURNED
    elif statuses == {PurchasePlanLine.REJECTED}:
        plan.status = PurchasePlan.REJECTED
    else:
        plan.status = PurchasePlan.APPROVED
        plan.approved_by = user
        plan.approval_date = timezone.now()
    plan.decision_comment = overall_comment
    plan.save()
    audit(user, AuditLog.APPROVAL, f"Purchase plan {plan.number} {plan.get_status_display()}",
          plan.number, new=overall_comment)
    msg = {PurchasePlan.APPROVED: "approved – raise the purchase orders in Odoo",
           PurchasePlan.RETURNED: "returned for revision",
           PurchasePlan.REJECTED: "rejected"}[plan.status]
    notify(plan.prepared_by, f"Purchase plan {plan.number} {msg}.", f"/plans/{plan.pk}/")
    return plan
